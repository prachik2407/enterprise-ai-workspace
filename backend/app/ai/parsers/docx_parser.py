from pathlib import Path

from docx import Document

from app.ai.parsers.base import BaseParser


class DOCXParser(BaseParser):
    """Parser for DOCX documents."""

    def parse(self, file_path: Path) -> str:
        """Extract text from paragraphs and tables in a DOCX file."""

        document = Document(str(file_path))

        text_parts: list[str] = []

        # Extract paragraphs
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Extract table content
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )

                if row_text:
                    text_parts.append(row_text)

        return "\n".join(text_parts)